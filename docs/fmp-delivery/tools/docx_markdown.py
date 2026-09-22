# docs/fmp-delivery/tools/docx_markdown.py
"""Render a small Markdown subset into an existing Word document's own formatting.

The delivery documents use direct formatting (no custom paragraph styles), so
new content is built by deep-copying a paragraph or table the document already
contains and replacing its runs. Supported inline syntax: ``**bold**``,
```code```, ``[label](url)``, ``[^N]`` (a reference to footnote id N of the
source document), and the team's review colours
``{yellow}…{/yellow}`` (also cyan, magenta, green). Colours are rendered only
after ``set_review(True)``; the submission build drops them. Supported blocks
are parsed by ``parse_blocks``.
"""

from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field

from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[\^\d+\]|\[[^\]]+\]\([^)]+\))")
FOOTNOTE = re.compile(r"\[\^(\d+)\]")
REVIEW_COLOURS = ("yellow", "cyan", "magenta", "green")
HIGHLIGHT = re.compile(r"\{(yellow|cyan|magenta|green)\}(.*?)\{/\1\}")
MARKER = re.compile(r"\{/?(?:yellow|cyan|magenta|green)\}")
_review = {"on": False}
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
CODE_FONT = "Consolas"
LINK_COLOR = "1b5fb4"
CONTENT_TAGS = {
    qn("w:r"),
    qn("w:sdt"),
    qn("w:hyperlink"),
    qn("w:bookmarkStart"),
    qn("w:bookmarkEnd"),
}


def set_review(on: bool) -> None:
    """Render the review colours (True) or drop them (False, the submission build)."""
    _review["on"] = on


def strip_markers(text: str) -> str:
    """Remove the review colour markers, keeping the text they wrap."""
    return MARKER.sub("", text)


def colour_segments(text: str) -> list[tuple[str, str | None]]:
    """Split text into (segment, colour) pieces; colour is None outside markers."""
    pieces: list[tuple[str, str | None]] = []
    position = 0
    for match in HIGHLIGHT.finditer(text):
        if match.start() > position:
            pieces.append((text[position : match.start()], None))
        pieces.append((match.group(2), match.group(1)))
        position = match.end()
    if position < len(text):
        pieces.append((text[position:], None))
    return pieces


@dataclass
class Block:
    """One block of the Markdown source."""

    kind: str  # h1 h2 h3 para bullet figcap tabcap table drawing image
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)


def parse_blocks(markdown: str) -> list[Block]:
    """Split Markdown into blocks: headings, tables, bullets, captions, images, paragraphs."""
    blocks: list[Block] = []
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line:
            continue
        if line.startswith("|"):
            rows = [line]
            while i < len(lines) and lines[i].startswith("|"):
                rows.append(lines[i].rstrip())
                i += 1
            cells = [_split_row(r) for r in rows if not re.fullmatch(r"\|[\s\-:|]+\|", r)]
            blocks.append(Block("table", rows=cells))
            continue
        blocks.append(_line_block(line))
    return blocks


def _split_row(row: str) -> list[str]:
    """Split a Markdown table row into stripped cells."""
    return [c.strip() for c in row.strip().strip("|").split("|")]


def _line_block(line: str) -> Block:
    """Classify a single non-table line."""
    for prefix, kind in (("### ", "h3"), ("## ", "h2"), ("# ", "h1"), ("- ", "bullet")):
        if line.startswith(prefix):
            return Block(kind, line[len(prefix) :])
    bare = strip_markers(line)
    drawing = re.fullmatch(r"\[\[DRAWING (\d+)\]\]", line)
    if drawing:
        return Block("drawing", drawing.group(1))
    image = re.fullmatch(r"\[\[IMAGE (.+)\]\]", line)
    if image:
        return Block("image", image.group(1))
    if re.match(r"Figure \d+\. ", bare):
        return Block("figcap", line)
    if re.match(r"Table [\dA-Z]+(\.\d+)?\. ", bare):
        return Block("tabcap", line)
    return Block("para", line)


def plain_text(text: str) -> str:
    """Strip inline Markdown and colour markers to the text a reader sees."""
    text = FOOTNOTE.sub("", LINK.sub(r"\1", strip_markers(text)))
    return text.replace("**", "").replace("`", "")


def base_rpr(paragraph_el) -> OxmlElement | None:
    """Return a copy of the rPr of the first run with text, without highlight."""
    for run in paragraph_el.iter(qn("w:r")):
        if run.find(qn("w:t")) is not None:
            rpr = run.find(qn("w:rPr"))
            rpr = copy.deepcopy(rpr) if rpr is not None else OxmlElement("w:rPr")
            for hl in rpr.findall(qn("w:highlight")):
                rpr.remove(hl)
            return rpr
    return OxmlElement("w:rPr")


def empty_copy(template_el):
    """Deep-copy a paragraph and drop its runs, keeping paragraph properties."""
    el = copy.deepcopy(template_el)
    for child in list(el):
        if child.tag in CONTENT_TAGS:
            el.remove(child)
    return el


# Child order of w:rPr in the OOXML schema; Word rejects runs whose properties are out of order.
RPR_ORDER = [
    "rStyle",
    "rFonts",
    "b",
    "bCs",
    "i",
    "iCs",
    "caps",
    "smallCaps",
    "strike",
    "dstrike",
    "outline",
    "shadow",
    "emboss",
    "imprint",
    "noProof",
    "snapToGrid",
    "vanish",
    "webHidden",
    "color",
    "spacing",
    "w",
    "kern",
    "position",
    "sz",
    "szCs",
    "highlight",
    "u",
    "effect",
    "bdr",
    "shd",
    "fitText",
    "vertAlign",
    "rtl",
    "cs",
    "em",
    "lang",
    "eastAsianLayout",
    "specVanish",
    "oMath",
]


def _set(rpr, tag: str, **attrs: str) -> None:
    """Set (replace) a single rPr child, inserting it at its schema position."""
    for old in rpr.findall(qn(tag)):
        rpr.remove(old)
    el = OxmlElement(tag)
    for key, value in attrs.items():
        el.set(qn(f"w:{key}"), value)
    rank = RPR_ORDER.index(tag.split(":")[1])
    for child in rpr:
        name = child.tag.split("}")[1]
        if name in RPR_ORDER and RPR_ORDER.index(name) > rank:
            child.addprevious(el)
            return
    rpr.append(el)


def make_run(
    text: str,
    rpr,
    *,
    bold: bool = False,
    code: bool = False,
    size: int | None = None,
    color: str | None = None,
    underline: bool = False,
    highlight: str | None = None,
):
    """Build a w:r element with a copy of ``rpr`` and the requested overrides."""
    run = OxmlElement("w:r")
    props = copy.deepcopy(rpr)
    if bold:
        _set(props, "w:b", val="1")
        _set(props, "w:bCs", val="1")
    if code:
        _set(props, "w:rFonts", ascii=CODE_FONT, hAnsi=CODE_FONT, cs=CODE_FONT, eastAsia=CODE_FONT)
    if color:
        _set(props, "w:color", val=color)
    if underline:
        _set(props, "w:u", val="single")
    if size:
        _set(props, "w:sz", val=str(size))
        _set(props, "w:szCs", val=str(size))
    if highlight:
        _set(props, "w:highlight", val=highlight)
    run.append(props)
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    run.append(t)
    return run


def add_inline(
    paragraph_el,
    text: str,
    rpr,
    part,
    *,
    size: int | None = None,
    code_size: int | None = None,
    bold_all: bool = False,
) -> None:
    """Append runs for ``text`` (inline Markdown and colour markers) to a paragraph element."""
    for segment, colour in colour_segments(text):
        hl = colour if _review["on"] else None
        for token in INLINE.split(segment):
            if not token:
                continue
            if token.startswith("**") and token.endswith("**"):
                run = make_run(token[2:-2], rpr, bold=True, size=size, highlight=hl)
            elif token.startswith("`") and token.endswith("`"):
                run = make_run(
                    token[1:-1], rpr, code=True, size=code_size or size, bold=bold_all, highlight=hl
                )
            elif FOOTNOTE.fullmatch(token):
                run = footnote_run(FOOTNOTE.fullmatch(token).group(1))
            elif LINK.fullmatch(token):
                label, url = LINK.fullmatch(token).groups()
                run = _hyperlink(
                    label, url, rpr, part, size=size, code_size=code_size, highlight=hl
                )
            else:
                run = make_run(token, rpr, size=size, bold=bold_all, highlight=hl)
            paragraph_el.append(run)


def footnote_run(footnote_id: str):
    """A superscript reference to an existing footnote, as the source document writes it."""
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    align = OxmlElement("w:vertAlign")
    align.set(qn("w:val"), "superscript")
    rpr.append(align)
    run.append(rpr)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), footnote_id)
    run.append(ref)
    return run


def _hyperlink(
    label: str,
    url: str,
    rpr,
    part,
    *,
    size: int | None,
    code_size: int | None,
    highlight: str | None = None,
):
    """Build an external w:hyperlink; a label wrapped in backticks keeps the code font."""
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    for token in INLINE.split(label):
        if not token:
            continue
        is_code = token.startswith("`") and token.endswith("`")
        link.append(
            make_run(
                token.strip("`"),
                rpr,
                code=is_code,
                size=(code_size if is_code else None) or size,
                color=LINK_COLOR,
                underline=True,
                highlight=highlight,
            )
        )
    return link


def paragraph(template_el, text: str, part, **kwargs):
    """New paragraph shaped like ``template_el`` carrying ``text``."""
    el = empty_copy(template_el)
    add_inline(el, text, base_rpr(template_el), part, **kwargs)
    return el


def strip_highlights(root) -> int:
    """Remove every w:highlight under ``root``; return how many were removed."""
    found = list(root.iter(qn("w:highlight")))
    for hl in found:
        hl.getparent().remove(hl)
    return len(found)


def table(
    template_tbl,
    rows: list[list[str]],
    part,
    *,
    total_width: int,
    font_size: int,
    code_size: int,
    weights: list[float] | None = None,
    highlight: str | None = None,
):
    """Build a table from ``template_tbl`` (header row + body row) with new rows.

    Column widths are proportional to ``weights`` or, when absent, to the longest
    plain-text cell of each column, with a floor so short columns stay legible.
    ``highlight`` colours every non-empty cell (review build only).
    """
    tbl = copy.deepcopy(template_tbl)
    tpl_rows = tbl.findall(qn("w:tr"))
    header_tr, body_tr = tpl_rows[0], tpl_rows[1]
    for tr in tpl_rows:
        tbl.remove(tr)
    style = tbl.find(qn("w:tblPr")).find(qn("w:tblStyle"))
    if style is not None:
        style.getparent().remove(style)

    ncols = max(len(r) for r in rows)
    widths = _widths(rows, ncols, total_width, weights)
    grid = tbl.find(qn("w:tblGrid"))
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(w))
        grid.append(col)

    for index, cells in enumerate(rows):
        tr = copy.deepcopy(header_tr if index == 0 else body_tr)
        for cant_split in tr.iter(qn("w:cantSplit")):
            cant_split.set(qn("w:val"), "1")  # a row never breaks across pages
        tc_template = tr.findall(qn("w:tc"))[0]
        for tc in tr.findall(qn("w:tc")):
            tr.remove(tc)
        for c, text in enumerate(cells + [""] * (ncols - len(cells))):
            tc = copy.deepcopy(tc_template)
            tc_pr = tc.find(qn("w:tcPr"))
            _set_cell_width(tc_pr, widths[c])
            p_template = tc.find(qn("w:p"))
            new_p = empty_copy(p_template)
            if highlight and text:
                text = f"{{{highlight}}}{text}{{/{highlight}}}"
            add_inline(new_p, text, base_rpr(p_template), part, size=font_size, code_size=code_size)
            tc.remove(p_template)
            tc.append(new_p)
            tr.append(tc)
        tbl.append(tr)
    return tbl


def _widths(rows, ncols, total, weights):
    """Column widths in dxa."""
    if weights is None:
        weights = []
        for c in range(ncols):
            body = max((len(plain_text(r[c])) if c < len(r) else 0) for r in rows[1:])
            header = len(plain_text(rows[0][c])) / 2 if c < len(rows[0]) else 0
            weights.append(max(8.0, min(float(max(body, header)), 60.0)))
    scale = total / sum(weights)
    widths = [int(w * scale) for w in weights]
    widths[0] += total - sum(widths)
    return widths


def _set_cell_width(tc_pr, width: int) -> None:
    """Set w:tcW on a cell's properties (first child, as the schema requires)."""
    for old in tc_pr.findall(qn("w:tcW")):
        tc_pr.remove(old)
    tcw = OxmlElement("w:tcW")
    tcw.set(qn("w:w"), str(width))
    tcw.set(qn("w:type"), "dxa")
    tc_pr.insert(0, tcw)


def add_picture(template_el, image_path: str, width_emu: int, document):
    """New picture paragraph shaped like ``template_el``."""
    el = empty_copy(template_el)
    para = Paragraph(el, document._body)
    para.add_run().add_picture(image_path, width=width_emu)
    return el


def body_text(document, stop_headings: tuple[str, ...]) -> list[tuple[str, str]]:
    """(current Heading 1, paragraph text) for every non-empty body paragraph."""
    out: list[tuple[str, str]] = []
    heading = ""
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style.name == "Heading 1":
            heading = text
        if heading in stop_headings:
            continue
        out.append((heading, text))
    return out
