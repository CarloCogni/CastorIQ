# docs/fmp-delivery/tools/rewrite_memory.py
"""Rebuild the body of the final memory from the Markdown section sources.

Reads ``delivery-docs/CastorIQ_Final_Memory_MAIN.docx`` (never modified), keeps
its cover page, and replaces everything from the Abstract onward with the
content of ``report-sections/00-*.md`` … ``11-*.md``, using the document's own
paragraph and table formatting as templates. Writes ``…_MAIN_v3.docx``.

Before saving it checks that every fraction and decimal in the evaluative
sections appears in an evaluation record, that no withdrawn figure or V2 term
survives, and that the body stays within the word limit.

    uv run --with python-docx python docs/fmp-delivery/tools/rewrite_memory.py
"""

from __future__ import annotations

import copy
import re
import sys
from pathlib import Path

import docx
import docx_markdown as md
from docx.oxml.ns import qn
from docx.shared import Emu

ROOT = Path(__file__).resolve().parents[3]
DELIVERY = ROOT / "docs/fmp-delivery/delivery-docs"
SOURCE = DELIVERY / "CastorIQ_Final_Memory_MAIN.docx"
TARGET = DELIVERY / "CastorIQ_Final_Memory_MAIN_v3.docx"
SECTIONS = ROOT / "docs/fmp-delivery/report-sections"
EVIDENCE = sorted((ROOT / "docs/evaluation").glob("*.md")) + sorted(
    (ROOT / "docs/evaluation/testing-log").glob("*.csv")
)
WORD_LIMIT = 5000
UNCOUNTED = ("Glossary", "References", "Appendices")
CHECKED_SECTIONS = ("00-", "05-", "06-")
TEXT_WIDTH_DXA = 8666  # the template table's width, i.e. the text column
EVIDENCE_MAP_WEIGHTS = [14, 26, 26, 22, 30]  # Table 5.1: result, harness, setup, record, artifacts
FORBIDDEN = (
    "REVIEW COPY",
    "Tier 0",
    "Tier 1,",
    "Tier 2,",
    "Tier 3",
    "qwen3:14b",
    "91.4",
    "78.3",
    "68.6",
    "19.2%",
    "0.41",
    "1,840",
    "Chen et al",
    "Ning et al",
    "Irani",
    "Autodesk / FMI",
)
# Figures legitimately absent from docs/evaluation, with where they come from.
ALLOWED_NUMBERS = {
    "0.10": "the requirement limit of the marginal RAV case; the record writes it as 0.10 in prose",
}
FRACTION = re.compile(r"(?<![\d.])(\d[\d,]*)\s*(?:/|of)\s*(\d[\d,]*)(?![\d.])")
SECTION_REF = re.compile(r"\b(?:Tables?|Sections?|Figures?)\s+\d+(?:\.\d+)*")
DECIMAL = re.compile(r"(?<![\d.])\d+\.\d+(?![\d.])")


class Templates:
    """Paragraph and table elements of the source document used as formatting templates."""

    def __init__(self, document: docx.document.Document) -> None:
        paragraphs = document.paragraphs
        self.label = self._find(paragraphs, lambda p: re.fullmatch(r"§ \d", p.text.strip()))
        self.h1 = self._find(paragraphs, lambda p: p.text.startswith("Introduction and problem"))
        self.h2 = self._find(paragraphs, lambda p: p.style.name == "Heading 2")
        self.body = self._find(paragraphs, lambda p: p.text.startswith("The AECO sector"))
        self.bullet = self._find(paragraphs, lambda p: p.text.startswith("Gap 1."))
        self.figcap = self._find(paragraphs, lambda p: p.text.startswith("Figure 2."))
        self.tabcap = self._find(paragraphs, lambda p: p.text.startswith("Table 1."))
        self.glossary = self._find(paragraphs, lambda p: p.text.startswith("AECO"))
        self.reference = self._find(paragraphs, lambda p: p.text.startswith("Autodesk"))
        self.pagebreak = self._find(
            paragraphs, lambda p: p._p.find(".//" + qn("w:br")) is not None and not p.text.strip()
        )
        self.table = copy.deepcopy(document.tables[1]._tbl)
        self.drawings = self._drawings(paragraphs)
        self.h3 = self._heading3()

    @staticmethod
    def _find(paragraphs, predicate):
        for p in paragraphs:
            if predicate(p):
                return copy.deepcopy(p._p)
        raise LookupError("template paragraph not found")

    @staticmethod
    def _drawings(paragraphs) -> dict[str, object]:
        """Map figure number → the (live) drawing paragraph that precedes its caption."""
        found = {}
        for before, p in zip(paragraphs, paragraphs[1:]):
            caption = re.match(r"Figure (\d+)\.", p.text)
            if caption and before._p.find(".//" + qn("w:drawing")) is not None:
                found[caption.group(1)] = before._p
        return found

    def _heading3(self):
        """Heading 3 does not occur in the source; derive it from Heading 2."""
        el = copy.deepcopy(self.h2)
        el.find(qn("w:pPr")).find(qn("w:pStyle")).set(qn("w:val"), "Heading3")
        spacing = el.find(qn("w:pPr")).find(qn("w:spacing"))
        spacing.set(qn("w:before"), "200")
        spacing.set(qn("w:after"), "80")
        rpr = el.find(qn("w:r")).find(qn("w:rPr"))
        for tag in ("w:sz", "w:szCs"):
            size = rpr.makeelement(qn(tag), {qn("w:val"): "22"})
            rpr.find(qn("w:rtl")).addprevious(size)
        return el


def keep_with_next(element) -> None:
    """Stop a heading, label or figure from being stranded at the bottom of a page."""
    ppr = element.find(qn("w:pPr"))
    if ppr.find(qn("w:keepNext")) is not None:
        return
    keep = ppr.makeelement(qn("w:keepNext"), {qn("w:val"): "1"})
    style = ppr.find(qn("w:pStyle"))
    if style is not None:
        style.addnext(keep)  # pStyle must stay the first child of pPr
    else:
        ppr.insert(0, keep)


def build(document, tpl: Templates) -> list:
    """Render every section file into a list of body elements."""
    part = document.part
    elements: list = []
    for path in sorted(SECTIONS.glob("[0-9][0-9]-*.md")):
        number = int(path.name[:2])
        blocks = md.parse_blocks(path.read_text(encoding="utf-8"))
        if 1 <= number <= 8:
            label = md.paragraph(tpl.label, f"§ {number}", part)
            keep_with_next(label)
            elements.append(label)
        for block in blocks:
            element = render(block, path.name, tpl, document)
            if block.kind in ("h1", "h2", "h3", "image", "drawing"):
                keep_with_next(element)
            elements.append(element)
        if number == 0:
            elements.append(copy.deepcopy(tpl.pagebreak))
    return elements


def render(block: md.Block, source: str, tpl: Templates, document):
    """One block → one body element."""
    part = document.part
    if block.kind == "h1":
        return md.paragraph(tpl.h1, block.text, part)
    if block.kind == "h2":
        return md.paragraph(tpl.h2, block.text, part)
    if block.kind == "h3":
        return md.paragraph(tpl.h3, block.text, part)
    if block.kind == "bullet":
        return md.paragraph(tpl.bullet, block.text, part, code_size=19)
    if block.kind == "figcap":
        return md.paragraph(tpl.figcap, block.text, part)
    if block.kind == "tabcap":
        return md.paragraph(tpl.tabcap, block.text, part, code_size=19)
    if block.kind == "drawing":
        return tpl.drawings[block.text]
    if block.kind == "image":
        return md.add_picture(
            tpl.drawings["2"], str(ROOT / block.text), Emu(TEXT_WIDTH_DXA * 635), document
        )
    if block.kind == "table":
        wide = len(block.rows[0]) >= 5
        evidence_map = block.rows[0][0] == "Result"
        return md.table(
            tpl.table,
            block.rows,
            part,
            total_width=TEXT_WIDTH_DXA,
            font_size=15 if wide else 17,
            code_size=13 if wide else 15,
            weights=EVIDENCE_MAP_WEIGHTS if evidence_map else None,
        )
    if source.startswith("09-") and block.text.startswith("**"):
        return glossary_entry(tpl.glossary, block.text)
    if source.startswith("10-"):
        return md.paragraph(tpl.reference, block.text, part)
    return md.paragraph(tpl.body, block.text, part, code_size=19)


def glossary_entry(template, text: str):
    """Bold term run in the glossary colour, then the definition run."""
    term, definition = re.fullmatch(r"\*\*(.+?)\*\*(.*)", text).groups()
    el = md.empty_copy(template)
    runs = [r for r in template.iter(qn("w:r")) if r.find(qn("w:t")) is not None]
    term_rpr = copy.deepcopy(runs[0].find(qn("w:rPr")))
    def_rpr = copy.deepcopy(runs[1].find(qn("w:rPr")))
    el.append(md.make_run(term, term_rpr))
    el.append(md.make_run(definition, def_rpr))
    return el


def replace_body(document, elements: list) -> None:
    """Keep the cover (everything before the Abstract heading), replace the rest."""
    body = document.element.body
    children = list(body.iterchildren())
    heading = next(p._p for p in document.paragraphs if p.text.strip() == "Abstract")
    abstract = children.index(heading)
    sect_pr = body.find(qn("w:sectPr"))
    before = {b.get(qn("r:embed")) for b in body.iter(qn("a:blip"))}
    for el in children[abstract:]:
        if el is not sect_pr:
            body.remove(el)
    for el in elements:
        sect_pr.addprevious(el)
    after = {b.get(qn("r:embed")) for b in body.iter(qn("a:blip"))}
    for r_id in before - after:  # images of figures that were redrawn
        document.part.drop_rel(r_id)


def check_numbers() -> list[str]:
    """Every fraction and decimal in the evaluative sections must occur in a record."""
    corpus = "\n".join(p.read_text(encoding="utf-8") for p in EVIDENCE)
    fractions = {f"{a}/{b}" for a, b in FRACTION.findall(corpus)}
    decimals = set(DECIMAL.findall(corpus))
    missing = []
    for path in sorted(SECTIONS.glob("[0-9][0-9]-*.md")):
        if not path.name.startswith(CHECKED_SECTIONS):
            continue
        text = md.plain_text(re.sub(r"\]\([^)]+\)", "]", path.read_text(encoding="utf-8")))
        text = re.sub(r"`[^`]*`", "", text)  # file names carry version numbers
        text = SECTION_REF.sub("", text)  # "Table 5.3", "Section 5.4" are not figures
        text = re.sub(r"(?m)^#+ .*$", "", text)  # nor are heading numbers
        for a, b in FRACTION.findall(text):
            if f"{a}/{b}" not in fractions:
                missing.append(f"{path.name}: {a} / {b}")
        for value in DECIMAL.findall(text):
            if value not in decimals and value not in ALLOWED_NUMBERS:
                missing.append(f"{path.name}: {value}")
    return sorted(set(missing))


def check_output(document) -> tuple[int, list[str]]:
    """Word count of the counted body, plus any forbidden string that survived."""
    counted = md.body_text(document, UNCOUNTED)
    words = sum(len(text.split()) for _, text in counted)
    full = "\n".join(p.text for p in document.paragraphs)
    full += "\n".join(c.text for t in document.tables for row in t.rows for c in row.cells)
    problems = [f"forbidden string survived: {s!r}" for s in FORBIDDEN if s in full]
    rsaa = [t for h, t in md.body_text(document, ()) if "RSAA" in t and h != "Glossary"]
    problems += [f"RSAA outside the glossary: {t[:60]}" for t in rsaa]
    highlights = sum(1 for _ in document.element.body.iter(qn("w:highlight")))
    if highlights:
        problems.append(f"{highlights} highlights left")
    if words > WORD_LIMIT:
        problems.append(f"body is {words} words, limit {WORD_LIMIT}")
    return words, problems


def main() -> int:
    """Build, check, save."""
    missing = check_numbers()
    for line in missing:
        print(f"UNSOURCED  {line}")
    document = docx.Document(str(SOURCE))
    templates = Templates(document)
    replace_body(document, build(document, templates))
    removed = md.strip_highlights(document.element.body)
    words, problems = check_output(document)
    for line in problems:
        print(f"PROBLEM    {line}")
    print(f"highlights removed from the cover: {removed}")
    print(
        f"counted body words (paragraphs, excluding {', '.join(UNCOUNTED)}; tables excluded): {words}"
    )
    if missing or problems:
        print("not saved")
        return 1
    document.save(str(TARGET))
    print(f"saved {TARGET.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
