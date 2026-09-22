# docs/fmp-delivery/tools/rewrite_appendices.py
"""Bring the appendices in line with the rewritten write path and add Appendix E.

Reads ``delivery-docs/CastorIQ_Final_Memory_APPENDICES.docx`` (never modified)
and writes ``…_APPENDICES_v3.docx``:

- Appendix A: the Modify paragraph and two captions describe the current
  models; Figure A.7 is replaced by the regenerated model graph.
- Appendix D: the stale "§2.6" pointer and the duplicated figure numbers
  D.3 and D.7 are fixed.
- Appendix E is appended from ``report-sections/appendix-e.md``; its table is
  the findings table of the dated record, so the two cannot drift apart.
- Appendix F is appended from ``report-sections/appendix-f.md``: where every
  figure of the memory's Section 5.2 lives in the repository. Its ``[[path]]``
  and ``[[path::symbol]]`` references become links at the tag
  (``code_links.py``); a missing file or symbol, or a linked file git does not
  track, stops the build.

With ``--review`` it also writes ``…_APPENDICES_v3_REVIEW.docx``: the review
note from ``report-sections/review-legend-appendices.md`` at the top, Appendix E
in cyan, and Word tracked changes against the original, verified the same way
as the memory's review copy.

    uv run --with python-docx python docs/fmp-delivery/tools/rewrite_appendices.py [--review]
"""

from __future__ import annotations

import copy
import re
import subprocess
import sys
from pathlib import Path

import code_links
import docx
import docx_markdown as md
import docx_redline as redline
from docx.oxml.ns import qn
from docx.shared import Emu

ROOT = Path(__file__).resolve().parents[3]
DELIVERY = ROOT / "docs/fmp-delivery/delivery-docs"
SOURCE = DELIVERY / "CastorIQ_Final_Memory_APPENDICES.docx"
TARGET = DELIVERY / "CastorIQ_Final_Memory_APPENDICES_v3.docx"
REVIEW_TARGET = DELIVERY / "CastorIQ_Final_Memory_APPENDICES_v3_REVIEW.docx"
REVIEW_NOTE = ROOT / "docs/fmp-delivery/report-sections/review-legend-appendices.md"
MAIN = DELIVERY / "CastorIQ_Final_Memory_MAIN.docx"
APPENDIX_E = ROOT / "docs/fmp-delivery/report-sections/appendix-e.md"
APPENDIX_F = ROOT / "docs/fmp-delivery/report-sections/appendix-f.md"
RECORD = ROOT / "docs/evaluation/2026-09-16-expert-testing-v3.md"
FIGURE_A7 = ROOT / "docs/fmp-delivery/figures/figA7-writeback-models.png"
TEXT_WIDTH_DXA = 8666
MAX_FIGURE_HEIGHT_EMU = int(8.2 * 914400)
FINDINGS_WEIGHTS = [8, 17, 14, 61]
EVIDENCE_WEIGHTS = [26, 74]  # Appendix F: what, where

A5_TEXT = (
    "The Modify path centres on ModificationProposal. Since the rewrite of 15 September 2026 it holds "
    "the request, the generated code, the target GlobalIds, the measured diff, the fingerprint of the file "
    "the code ran against, the path of the reviewed scratch copy, the model that wrote the blind "
    "explanation, whether the document check was skipped, when flagged rows were acknowledged, and the "
    "Guardian result. The tier, operation, intent and filter fields are kept for the history of the first, "
    "tiered design. Every applied proposal is linked one-to-one with a GitCommit record that stores the "
    "commit hash, affected entity counts, diff data, and rollback flag. Conflict records the outputs of the "
    "document cross-check between document chunks and IFC entities, with severity, resolution notes, and a "
    "suggested fix. ScanRun groups conflict detections into batches so that per-run statistics stay "
    "auditable. FailureRecord persists failed write-path attempts with a deterministic failure class, the "
    "failing phase, a diagnosis, and whether a retry can help."
)
REPLACEMENTS = {
    "The Modify path centres on ModificationProposal": A5_TEXT,
    "Figure A.7.": (
        "Figure A.7. Modify and Conflicts: proposal, commit, conflict, and scan-run models, "
        "regenerated after the rewrite of 15 September 2026."
    ),
    "Figure A.8.": (
        "Figure A.8. FailureRecord: structured diagnostics for failed write-path attempts; "
        "the tier and intent fields belong to the first design."
    ),
}


def set_text(paragraph, text: str) -> None:
    """Replace a paragraph's text, keeping the formatting of its first text run."""
    runs = [r for r in paragraph._p.iter(qn("w:r")) if r.find(qn("w:t")) is not None]
    runs[0].find(qn("w:t")).text = text
    for run in runs[1:]:
        run.getparent().remove(run)


def fix_appendix_a(document) -> list[str]:
    """Rewrite the stale Modify paragraph and captions; swap Figure A.7."""
    done = []
    paragraphs = document.paragraphs
    for index, p in enumerate(paragraphs):
        for prefix, text in REPLACEMENTS.items():
            if p.text.startswith(prefix):
                set_text(p, text)
                done.append(prefix)
        if p.text.startswith("Figure A.7."):
            swap_picture(document, paragraphs[index - 1]._p, FIGURE_A7)
            done.append("Figure A.7 image")
    return done


def swap_picture(document, drawing_p, image: Path) -> None:
    """Replace the picture in ``drawing_p`` with ``image`` at the same width, capped in height."""
    old_blip = drawing_p.find(".//" + qn("a:blip"))
    old_rid = old_blip.get(qn("r:embed"))
    width = int(drawing_p.find(".//" + qn("wp:extent")).get("cx"))
    new_p = md.empty_copy(drawing_p)
    run = docx.text.paragraph.Paragraph(new_p, document._body).add_run()
    shape = run.add_picture(str(image), width=Emu(width))
    if shape.height > MAX_FIGURE_HEIGHT_EMU:
        ratio = MAX_FIGURE_HEIGHT_EMU / shape.height
        shape.height = Emu(MAX_FIGURE_HEIGHT_EMU)
        shape.width = Emu(int(width * ratio))
    drawing_p.addprevious(new_p)
    drawing_p.getparent().remove(drawing_p)
    document.part.drop_rel(old_rid)


def fix_appendix_d(document) -> list[str]:
    """Fix the section pointer and give the two doubled figure numbers a and b."""
    done = []
    seen: dict[str, int] = {}
    captions = [p for p in document.paragraphs if re.match(r"Figure D\.\d+\.", p.text)]
    counts: dict[str, int] = {}
    for p in captions:
        number = p.text.split(".", 2)[1]
        counts[number] = counts.get(number, 0) + 1
    for p in captions:
        number = p.text.split(".", 2)[1]
        if counts[number] < 2:
            continue
        seen[number] = seen.get(number, 0) + 1
        suffix = "ab"[seen[number] - 1]
        set_text(p, p.text.replace(f"Figure D.{number}.", f"Figure D.{number}{suffix}.", 1))
        done.append(f"Figure D.{number}{suffix}")
    for p in document.paragraphs:
        if "§2.6" in p.text:
            set_text(p, p.text.replace("lives in §2.6", "is in Section 4.4 of the memory"))
            done.append("§2.6 pointer")
    return done


def record_table() -> list[list[str]]:
    """The findings table of the dated record."""
    text = RECORD.read_text(encoding="utf-8")
    section = text.split("## The fourteen rows", 1)[1].split("\n## ", 1)[0]
    return next(b for b in md.parse_blocks(section) if b.kind == "table").rows


def appendix(
    document,
    markdown: str,
    *,
    record_rows: list[list[str]] | None = None,
    table_weights: list[float] | None = None,
    table_highlight: str | None = None,
) -> list:
    """Build one appended appendix in the document's own appendix formatting.

    ``[[RECORD TABLE]]`` stands for ``record_rows``; other tables are rendered
    from the Markdown. Bullets become indented body paragraphs, because the
    appendices define no list numbering.
    """
    paragraphs = document.paragraphs
    page_break = next(
        p._p for p in paragraphs if p._p.find(".//" + qn("w:pageBreakBefore")) is not None
    )
    title = next(p._p for p in paragraphs if p.text.startswith("Appendix D:"))
    heading = next(p._p for p in paragraphs if p.text.startswith("D.1 Overview"))
    body = next(
        p._p for p in paragraphs if p.text.startswith("This appendix reproduces the client-facing")
    )
    caption = next(p._p for p in paragraphs if p.text.startswith("Figure D.1."))
    table_template = docx.Document(str(MAIN)).tables[1]._tbl
    part = document.part

    def table(rows: list[list[str]]):
        return md.table(
            table_template,
            rows,
            part,
            total_width=TEXT_WIDTH_DXA,
            font_size=15,
            code_size=13,
            weights=table_weights,
            highlight=table_highlight,  # rendered in the review copy only
        )

    elements = [copy.deepcopy(page_break)]
    for block in md.parse_blocks(markdown):
        if block.kind == "h1":
            elements.append(md.paragraph(title, block.text, part))
        elif block.kind == "h2":
            elements.append(md.paragraph(heading, block.text, part))
        elif block.kind == "tabcap":
            elements.append(md.paragraph(caption, block.text, part, code_size=16))
        elif block.kind == "table":
            elements.append(table(block.rows))
        elif block.text == "[[RECORD TABLE]]":
            elements.append(table(record_rows or []))
        elif block.kind == "bullet":
            elements.append(bullet(body, block.text, part))
        else:
            elements.append(md.paragraph(body, block.text, part, code_size=19))
    return elements


def bullet(template, text: str, part):
    """A body paragraph with a hanging bullet."""
    el = md.paragraph(template, "•\u00a0\u00a0" + text, part, code_size=19)
    ppr = el.find(qn("w:pPr"))
    ind = ppr.makeelement(qn("w:ind"), {qn("w:left"): "360", qn("w:hanging"): "240"})
    jc = ppr.find(qn("w:jc"))
    if jc is None:
        jc = ppr.makeelement(qn("w:jc"), {})
        ppr.append(jc)
    jc.set(qn("w:val"), "left")  # justified bullets stretch around long code paths
    jc.addprevious(ind)
    return el


def appendix_e(document) -> list:
    """Appendix E: the testing log of the rewritten write path, table from the record."""
    return appendix(
        document,
        APPENDIX_E.read_text(encoding="utf-8"),
        record_rows=record_table(),
        table_weights=FINDINGS_WEIGHTS,
        table_highlight="cyan",
    )


def appendix_f(document) -> tuple[list, list[str]]:
    """Appendix F: where every §5.2 figure lives; returns the elements and linked paths."""
    markdown = code_links.resolve(APPENDIX_F.read_text(encoding="utf-8"), ROOT)
    linked = re.findall(r"/(?:blob|tree)/" + code_links.TAG + r"/([^)#]+)", markdown)
    return appendix(document, markdown, table_weights=EVIDENCE_WEIGHTS), linked


def untracked(paths: list[str]) -> list[str]:
    """Linked paths git does not track (they would 404 at the tag)."""
    tracked = subprocess.run(
        ["git", "ls-files"], cwd=ROOT, capture_output=True, text=True, check=True
    ).stdout.splitlines()
    files = set(tracked)
    return sorted(
        {p for p in paths if p not in files and not any(t.startswith(p + "/") for t in tracked)}
    )


def review_note(document) -> list:
    """The review callout for the top of the appendices, in the memory's callout format."""
    paragraphs = docx.Document(str(MAIN)).paragraphs
    label = next(p._p for p in paragraphs if "REVIEW COPY" in p.text)
    legend = next(p._p for p in paragraphs if p.text.strip().startswith("Yellow ="))
    elements = []
    for block in md.parse_blocks(REVIEW_NOTE.read_text(encoding="utf-8")):
        template = label if block.kind == "h2" else legend
        elements.append(md.paragraph(template, "  " + block.text, document.part, code_size=15))
    return elements


def build_document(review: bool):
    """The fixed appendices with Appendix E; in review mode also the note and colours."""
    md.set_review(review)
    document = docx.Document(str(SOURCE))
    done = fix_appendix_a(document) + fix_appendix_d(document)
    body = document.element.body
    sect_pr = body.find(qn("w:sectPr"))
    f_elements, linked = appendix_f(document)
    for element in appendix_e(document) + f_elements:
        sect_pr.addprevious(element)
    missing = untracked(linked)
    if missing:
        raise code_links.CodeLinkError(f"linked but not tracked by git: {', '.join(missing)}")
    notes = review_note(document) if review else []
    first = next(body.iterchildren())
    for element in notes:
        first.addprevious(element)
    return document, done, redline.lines_of(notes)


def build_review(clean) -> int:
    """Write the review copy with tracked changes; verify the redline."""
    document, _, review_only = build_document(review=True)
    original = docx.Document(str(SOURCE))
    counts = redline.mark(original, document)
    problems = redline.check(document, clean, original, review_only)
    print(f"review: {counts}")
    for line in problems:
        print(f"PROBLEM    {line}")
    if problems:
        print("review copy not saved")
        return 1
    document.save(str(REVIEW_TARGET))
    print(f"saved {REVIEW_TARGET.relative_to(ROOT)}")
    return 0


def main(argv: list[str]) -> int:
    """Apply the fixes, append Appendix E, check, save; ``--review`` adds the review copy."""
    document, done, _ = build_document(review=False)
    for item in done:
        print(f"fixed      {item}")
    full = "\n".join(p.text for p in document.paragraphs)
    problems = [s for s in ("§2.6", "tier assignment", "tier calibration", "Tier-1") if s in full]
    doubled = [
        n
        for n in set(re.findall(r"^Figure (D\.\d+)\.", full, re.M))
        if len(re.findall(rf"^Figure {re.escape(n)}\.", full, re.M)) > 1
    ]
    for problem in problems + doubled:
        print(f"PROBLEM    {problem}")
    if problems or doubled or len(done) < 8:
        print("not saved")
        return 1
    document.save(str(TARGET))
    print(f"saved {TARGET.relative_to(ROOT)}")
    if "--review" not in argv:
        return 0
    return build_review(document)


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
